import argparse
import csv
import hashlib
import json
import math
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def printable(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    return value


def minute(value):
    if isinstance(value, (time, datetime)):
        return value.hour * 60 + value.minute
    text = str(value).strip()
    extra = 1440 if text.endswith("+1") else 0
    h, m = text.replace("+1", "").split(":")[:2]
    return 60 * int(h) + int(m) + extra


def numeric(values, location):
    bad = [i for i, x in enumerate(np.asarray(values, dtype=object).flat)
           if not isinstance(x, (int, float)) or isinstance(x, bool) or not math.isfinite(x)]
    if bad:
        raise ValueError(f"Nonfinite/missing/nonnumeric source cells in {location}: {bad[:10]}")
    return np.asarray(values, dtype=float)


def statistics(values):
    x = np.asarray(values, dtype=float)
    flat = x.ravel()
    delta = np.abs(np.diff(x, axis=-1)).ravel()
    q1, q3 = np.quantile(delta, [0.25, 0.75])
    threshold = float(q3 + 3 * (q3 - q1))
    return {
        "count": int(flat.size), "missing_or_nonfinite": int((~np.isfinite(flat)).sum()),
        "negative": int((flat < 0).sum()), "zero": int((flat == 0).sum()),
        "min": float(flat.min()), "max": float(flat.max()),
        "mean": float(flat.mean()), "median": float(np.median(flat)),
        "p01": float(np.quantile(flat, .01)), "p99": float(np.quantile(flat, .99)),
        "maximum_adjacent_change": float(delta.max()),
        "adjacent_change_screen_threshold": threshold,
        "adjacent_changes_above_screen": int((delta > threshold).sum()),
        "screen_only_no_values_removed": True,
        "duplicate_daily_profiles": int(len(x) - len(np.unique(x, axis=0))) if x.ndim == 2 else None,
    }


def inventory(path, template=False):
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=False)
    sheets = []
    for sheet in workbook:
        item = {
            "name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column,
            "merged_cells": [str(r) for r in sheet.merged_cells.ranges],
            "header": [printable(c.value) for c in sheet[1]],
            "formula_cells": [c.coordinate for row in sheet for c in row if c.data_type == "f"],
            "populated_cells": sum(c.value not in (None, "") for row in sheet for c in row),
            "hidden": sheet.sheet_state,
        }
        if template:
            item["fixed_cells"] = {
                c.coordinate: printable(c.value) for row in sheet for c in row
                if c.value not in (None, "")
            }
        sheets.append(item)
    workbook.close()
    return {"bytes": path.stat().st_size, "sha256": digest(path), "sheets": sheets}


def read_daily(sheet):
    rows = list(sheet.values)
    ends = [minute(v) for v in rows[0][1:]]
    if ends != list(range(10, 1441, 10)):
        raise ValueError(f"Unexpected 10-minute grid in {sheet.title}")
    days = [v[0].date() if isinstance(v[0], datetime) else v[0] for v in rows[1:]]
    expected = [date(2025, 1, 1) + timedelta(days=i) for i in range(365)]
    if days != expected:
        raise ValueError(f"Incomplete/duplicate/unordered dates in {sheet.title}")
    values = numeric([v[1:] for v in rows[1:]], sheet.title)
    if values.shape != (365, 144):
        raise ValueError(f"Unexpected numeric shape in {sheet.title}: {values.shape}")
    return days, values


def write_csv(path, columns, rows):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(columns)
        writer.writerows(rows)


def run(config_path):
    root = Path(__file__).resolve().parents[1]
    config_path = root / config_path
    config = json.loads(config_path.read_text())
    raw = root / config["raw_dir"]
    templates = root / config["template_dir"]
    output = root / config["processed_dir"]
    audit_dir = root / "results/audit"
    output.mkdir(parents=True, exist_ok=True)
    audit_dir.mkdir(parents=True, exist_ok=True)
    all_inputs = sorted(raw.glob("*.xlsx")) + sorted(templates.glob("*.xlsx"))
    before = {str(p.relative_to(root)): inventory(p, p.parent == templates) for p in all_inputs}
    books = {p.name: openpyxl.load_workbook(p, data_only=True, read_only=True) for p in raw.glob("*.xlsx")}
    q1rows = list(books["附件1.xlsx"].active.values)
    if [minute(row[0]) for row in q1rows[1:]] != list(range(10, 1441, 10)):
        raise ValueError("Attachment 1 does not have exactly 144 right endpoints")
    q1 = numeric([row[1:] for row in q1rows[1:]], "附件1")
    if q1.shape != (144, 3):
        raise ValueError(f"Unexpected Attachment 1 shape: {q1.shape}")
    days, load = read_daily(books["附件2.xlsx"]["小区负载"])
    pvdays, pv = read_daily(books["附件2.xlsx"]["光伏发电实际功率"])
    pricedays, price = read_daily(books["附件4.xlsx"].active)
    assert days == pvdays == pricedays
    q1cols = ["interval_index", "start_minute", "end_minute", "price", "load_kw", "pv_kw", "load_kwh", "pv_kwh"]
    write_csv(output / "q1.csv", q1cols,
              ([i, i * 10, (i + 1) * 10, p, l, v, l / 6, v / 6]
               for i, (p, l, v) in enumerate(q1)))
    point_mean = .5 * (q1[:, 1:] + np.roll(q1[:, 1:], 1, axis=0))
    write_csv(output / "q1_instantaneous_sensitivity.csv", q1cols,
              ([i, i * 10, (i + 1) * 10, q1[i, 0], l, v, l / 6, v / 6]
               for i, (l, v) in enumerate(point_mean)))
    actual_columns = ["date", "interval_index", "interval_start", "interval_end", "load_kw", "pv_kw",
                      "price_variable", "load_kwh", "pv_kwh", "available_at"]
    def actual_rows():
        for d, day in enumerate(days):
            day_start = datetime.combine(day, time())
            for t in range(144):
                start = day_start + timedelta(minutes=10 * t)
                end = start + timedelta(minutes=10)
                yield [day.isoformat(), t, start.isoformat(), end.isoformat(), load[d, t], pv[d, t],
                       price[d, t], load[d, t] / 6, pv[d, t] / 6, end.isoformat()]
    write_csv(output / "actuals.csv", actual_columns, actual_rows())
    issue_times, forecasts, forecast_rows = [], [], []
    current_date, filled = None, 0
    for row_number, row in enumerate(books["附件3.xlsx"].active.iter_rows(min_row=2, values_only=True), 2):
        label, hour, *values = row
        if label not in (None, ""):
            current_date = datetime.strptime(str(label), "%Y-%m-%d").date()
        else:
            filled += 1
        if current_date is None:
            raise ValueError(f"Forecast date missing before first nonempty date at row {row_number}")
        issue = datetime.combine(current_date, time()) + timedelta(minutes=minute(hour))
        value = numeric(values, f"附件3 row {row_number}")
        if value.size != 24:
            raise ValueError(f"Forecast horizon not 24 at row {row_number}")
        issue_times.append(issue)
        forecasts.append(value)
        for lead, v in enumerate(value, 1):
            forecast_rows.append([issue.isoformat(), issue.isoformat(), (issue + timedelta(hours=lead)).isoformat(),
                                  lead, float(v), row_number, lead + 2])
    expected_issues = [datetime(2025, 1, 1) + timedelta(hours=6 * i) for i in range(365 * 4)]
    if issue_times != expected_issues:
        raise ValueError("Forecast release times are incomplete, duplicate or out of order")
    forecasts = np.asarray(forecasts)
    write_csv(output / "forecasts.csv",
              ["issue_time", "available_at", "target_time", "lead_time", "value", "source_row", "source_column"], forecast_rows)
    np.savez_compressed(output / "arrays.npz", dates=np.array([x.isoformat() for x in days]),
                        load_kw=load, pv_kw=pv, price_variable=price, price_fixed=q1[:, 0],
                        q1_load_kw=q1[:, 1], q1_pv_kw=q1[:, 2],
                        forecast_values=forecasts,
                        forecast_issue_times=np.array([x.isoformat() for x in issue_times]))
    stats = {
        "q1_price": statistics(q1[:, 0]), "q1_load_kw": statistics(q1[:, 1]), "q1_pv_kw": statistics(q1[:, 2]),
        "actual_load_kw": statistics(load), "actual_pv_kw": statistics(pv),
        "price_variable": statistics(price), "forecast_pv_kw": statistics(forecasts),
    }
    relationships = {
        "q1_load_vs_annual_slot_mean_max_abs_kw": float(np.max(np.abs(q1[:, 1] - load.mean(axis=0)))),
        "q1_pv_vs_annual_slot_mean_max_abs_kw": float(np.max(np.abs(q1[:, 2] - pv.mean(axis=0)))),
        "q1_price_vs_annual_slot_mean_max_abs": float(np.max(np.abs(q1[:, 0] - price.mean(axis=0)))),
        "interpretation": "Observed numeric similarity does not prove how the organizer created the inputs. Do not use Attachment 1 load/PV to train Q2-4 or infer that future variable prices are published in advance.",
    }
    checks = {
        "all_raw_numeric_cells_finite": all(v["missing_or_nonfinite"] == 0 for v in stats.values()),
        "all_raw_values_nonnegative": all(v["negative"] == 0 for v in stats.values()),
        "daily_grid_144_intervals_145_state_endpoints": True,
        "date_coverage_365_days": len(days) == 365,
        "evaluation_334_days": sum(d >= date(2025, 2, 1) for d in days) == 334,
        "evaluation_48096_intervals": sum(d >= date(2025, 2, 1) for d in days) * 144 == 48096,
        "forecast_1460_issues_35040_values": forecasts.shape == (1460, 24),
        "forecast_unique_issue_lead_pairs": len({(r[0], r[3]) for r in forecast_rows}) == 35040,
        "forecast_available_at_equals_issue_time": all(r[0] == r[1] for r in forecast_rows),
        "forecast_target_after_release": all(1 <= r[3] <= 24 for r in forecast_rows),
        "5000_kw_for_10_minutes_is_833_333_kwh": math.isclose(5000 / 6, 833.3333333333334),
        "raw_and_templates_hash_unchanged": all(digest(root / p) == info["sha256"] for p, info in before.items()),
    }
    readback = pd.read_csv(output / "actuals.csv")
    q1back = pd.read_csv(output / "q1.csv")
    forecastsback = pd.read_csv(output / "forecasts.csv")
    residuals = {
        "q1_load_energy_kwh": float(abs(q1back.load_kwh.sum() - q1[:, 1].sum() / 6)),
        "q1_pv_energy_kwh": float(abs(q1back.pv_kwh.sum() - q1[:, 2].sum() / 6)),
        "annual_load_energy_kwh": float(abs(readback.load_kwh.sum() - load.sum() / 6)),
        "annual_pv_energy_kwh": float(abs(readback.pv_kwh.sum() - pv.sum() / 6)),
        "csv_power_roundtrip_max_abs_kw": float(max(np.max(abs(readback.load_kw.to_numpy().reshape(load.shape) - load)),
                                                     np.max(abs(readback.pv_kw.to_numpy().reshape(pv.shape) - pv)))),
        "forecast_roundtrip_max_abs_kw": float(np.max(abs(forecastsback.value.to_numpy().reshape(forecasts.shape) - forecasts))),
        "point_sensitivity_daily_load_energy_difference_kwh": float(abs(point_mean[:, 0].sum() / 6 - q1[:, 1].sum() / 6)),
        "point_sensitivity_daily_pv_energy_difference_kwh": float(abs(point_mean[:, 1].sum() / 6 - q1[:, 2].sum() / 6)),
    }
    checks["csv_readback_energy_and_power_conserved"] = all(v <= 1e-7 for v in residuals.values())
    starts = pd.to_datetime(readback.interval_start)
    ends = pd.to_datetime(readback.interval_end)
    checks["all_intervals_ten_minutes"] = bool(((ends - starts).dt.total_seconds() == 600).all())
    checks["no_duplicate_or_missing_cross_day_intervals"] = bool(starts.is_unique and (starts.iloc[1:].to_numpy() == ends.iloc[:-1].to_numpy()).all())
    issues = [
        {"id": "DQ001", "kind": "source_semantics_ambiguity", "status": "DECLARED_ASSUMPTION",
         "description": "Raw labels are 00:10 through 24:00; no statement specifies interval averages versus point measurements. Main path treats power as preceding-interval averages. Q1 point-linear periodic-boundary sensitivity is provided separately."},
        {"id": "DQ002", "kind": "template_time_labels", "status": "REQUIRES_DOCUMENTED_EXPORT_MAPPING",
         "description": "All five templates start at 00:10-00:20 and omit same-day 00:00-00:10. result1 ends at 00:00+1-00:10+1; other wide sheets end at ambiguous 00:00-00:10+1. Exact label-preserving export cannot simultaneously certify canonical same-day coverage without a disclosed mapping."},
        {"id": "DQ003", "kind": "template_placeholder_blocks", "status": "EXPAND_IN_OUTPUT_COPIES",
         "description": "Battery and emergency templates contain example date blocks and ellipses, not all 334 days. Expand output copies while preserving table columns, sheet names and formats."},
        {"id": "DQ004", "kind": "forecast_time_semantics", "status": "DECLARED_ASSUMPTION",
         "description": "Forecast columns are hour-ahead values at leads 1..24; there is no lead-0 value. Retain every release separately. Point-to-interval reconstruction must declare its causal lead-0 boundary and retain energy/nonnegative checks."},
        {"id": "DQ005", "kind": "price_information_availability", "status": "UNRESOLVED_IN_STATEMENT",
         "description": "Attachment 4 gives delivery-slot prices with no publication timestamps. Actual prices must not automatically become known day-ahead prices."},
        {"id": "DQ006", "kind": "cross_source_information_boundary", "status": "OFFLINE_OBSERVATION_ONLY",
         "description": "Attachment 1 load and price profiles are numerically very close to annual per-slot averages of Attachments 2 and 4. This is not proof of generation provenance and does not authorize using Attachment 1 load/PV in Q2-4 forecasting or knowing future variable prices."},
    ]
    totals = {"q1_load_kwh": float(q1[:, 1].sum() / 6), "q1_pv_kwh": float(q1[:, 2].sum() / 6),
              "annual_load_kwh": float(load.sum() / 6), "annual_pv_kwh": float(pv.sum() / 6),
              "evaluation_load_kwh": float(load[31:].sum() / 6), "evaluation_pv_kwh": float(pv[31:].sum() / 6)}
    report = {"run_id": "data-audit-" + datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ"),
              "config": str(config_path.relative_to(root)), "code_sha256": digest(Path(__file__)),
              "status": "PASS_WITH_DECLARED_SEMANTIC_ISSUES" if all(checks.values()) else "FAIL",
              "inputs": before, "statistics": stats, "checks": checks, "residuals": residuals,
              "totals": totals, "issues": issues, "cross_source_relationships": relationships,
              "forecast_date_cells_forward_filled": filled,
              "forecast_first_issue": issue_times[0].isoformat(), "forecast_last_issue": issue_times[-1].isoformat(),
              "forecast_last_target": forecast_rows[-1][2], "timezone": "naive_problem_local_calendar_no_DST",
              "outputs": {str(p.relative_to(root)): {"bytes": p.stat().st_size, "sha256": digest(p)}
                          for p in sorted(output.iterdir()) if p.is_file()},
              "environment": {"numpy": np.__version__, "pandas": pd.__version__, "openpyxl": openpyxl.__version__}}
    (audit_dir / "data_audit.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n")
    write_csv(audit_dir / "data_hashes.csv", ["path", "bytes", "sha256"],
              ([p, info["bytes"], info["sha256"]] for p, info in before.items()))
    daily = pd.DataFrame({"date": [d.isoformat() for d in days], "load_kwh": load.sum(axis=1) / 6,
                          "pv_kwh": pv.sum(axis=1) / 6, "price_min": price.min(axis=1),
                          "price_max": price.max(axis=1), "price_mean": price.mean(axis=1)})
    daily.to_csv(audit_dir / "data_daily_summary.csv", index=False)
    write_documents(root, report)
    for workbook in books.values():
        workbook.close()
    print(json.dumps({"status": report["status"], "checks_passed": sum(checks.values()), "checks_total": len(checks),
                      "totals": totals, "statistics": stats, "issues": [x["id"] for x in issues]}, ensure_ascii=False, indent=2))
    if not all(checks.values()):
        raise SystemExit(1)


def write_documents(root, report):
    dictionary = """# C 题数据字典

来源：题面第 2–3 页、归档后附件 1–4、附件 5 五个空白模板。由 `src/data_audit.py` 复现。

| 原件 | 实际表及布局 | 数值单位 | 题定用途 |
| --- | --- | --- | --- |
| data/raw/2026C/附件1.xlsx | Sheet1；144 行，时间/电价/小区负载/光伏发电预测功率 | 元/kWh；kW；kW | Q1 全部输入，Q2/Q3 固定每日价格 |
| data/raw/2026C/附件2.xlsx | 小区负载、光伏发电实际功率；各 365 日×144 时段 | kW | Q2–Q4 历史观测与仿真真值 |
| data/raw/2026C/附件3.xlsx | Sheet1；1460 发布×24 预报，另有日期与预报时刻 | kW | 仅 Q3/Q4-3 的已发布预报 |
| data/raw/2026C/附件4.xlsx | Sheet1；365 日×144 时段 | 元/kWh | Q4 交付时段的真实结算价格 |

原始日期覆盖 2025-01-01 至 2025-12-31；Q2–Q4 正式评估 2025-02-01 至 12-31，共 334 日、48096 个时段。全部采用题目当地无时区日历，不套用美国 DST。

## 时间和功率语义

附件 1/2/4 实际只有 144 个标签：00:10、00:20、……、23:50、0:00+1。最后标签等于当天 24:00，不是第 145 个供电区间。内部采用 144 个左闭右开时段 [00:00,00:10)、……、[23:50,24:00)，储能采用 145 个状态端点。

题面仅称“功率”，未明确瞬时值还是区间平均值。**主建模假设**：标签 t 对应此前 10 分钟的平均功率；区间电量 = 功率/6。价格同样归属此前交付区间。原件没有被改写。`q1_instantaneous_sensitivity.csv` 为端点瞬时功率的独立解释：相邻端点线性积分，以周期条件补 P(00:00)=P(24:00)，电价仍按交付区间计。周期补边是敏感性假设，不是原始观测。

## 派生文件

| 文件/字段 | 类型、单位与含义 |
| --- | --- |
| q1.csv / interval_index | 整数 0–143；当日唯一时段 |
| q1.csv / start_minute,end_minute | 相对当天 0 点的分钟；0–1430、10–1440 |
| q1.csv / price | 元/kWh；当天固定交付价格 |
| q1.csv / load_kw,pv_kw | kW；主假设下区间平均功率 |
| q1.csv / load_kwh,pv_kwh | kWh；对应功率除以 6 |
| actuals.csv / date,interval_index | 日期+整数联合主键；52560 行 |
| actuals.csv / interval_start,interval_end | ISO 当地时间；相邻区间端点衔接 |
| actuals.csv / load_kw,pv_kw,price_variable,load_kwh,pv_kwh | 真实数据及物理单位同上；仅仿真环境持有未来记录 |
| actuals.csv / available_at | 保守的已完成区间数据可用时点，等于 interval_end；并非官方公布时间。实时控制若使用当前区间观测须另声明执行信息假设 |
| forecasts.csv / issue_time,available_at | ISO 当地时间；预报发布时刻及假定即时可用时刻，二者相同 |
| forecasts.csv / target_time | issue_time+lead_time 小时；保留跨年目标 |
| forecasts.csv / lead_time | 整数小时，1–24，绝非 0–23 |
| forecasts.csv / value | kW；对应整点预报，未擅自插值为 10 分钟 |
| forecasts.csv / source_row,source_column | Excel 1 基坐标，保留值级来源 |
| arrays.npz | 同内容的数值缓存；数组日期/发布顺序与 CSV 一致；不得给计划器传可任意索引的全年真值 |

附件 3 日期空白为同一日期的延续：按首次非空日期向下填充 1095 个日期标签；没有填补任何缺失功率。每个 (issue_time,lead_time) 保留唯一记录，不能按 target_time 去重。末次发布 2025-12-31 18:00，末目标为 2026-01-01 18:00，这些跨年预报真实存在，不能当成额外的 2025 年观测。预报 lead 0 未给出，后续重采样要明确因果边界。
"""
    (root / "context/DATA_DICTIONARY.md").write_text(dictionary)
    stats = report["statistics"]
    table = "\n".join(f"| {name} | {v['count']} | {v['min']:.6g} | {v['max']:.6g} | {v['negative']} | {v['missing_or_nonfinite']} | {v['adjacent_changes_above_screen']} |"
                      for name, v in stats.items())
    quality = f"""# C 题数据质量

真实运行 `{report['run_id']}`：**{report['status']}**。机器检查 {sum(report['checks'].values())}/{len(report['checks'])} 通过；详细证据在 `results/audit/data_audit.json`，哈希清单为 `results/audit/data_hashes.csv`。本脚本未写入任何原件或模板。

## 全量数值审计

| 变量 | 个数 | 最小值 | 最大值 | 负值 | 缺失/非有限 | 相邻跳变筛查数 |
| --- | --- | --- | --- | --- | --- | --- |
{table}

跳变仅为**离线筛查**：同一日/同一发布的相邻差绝对值超过其全样本 Q3+3×IQR。它不是物理异常判决，也未用于在线训练、阈值校准、删除或修正数据。光伏日出日落、调价和真实负载变化可自然触发。原始数值全部保留；详细均值、分位数、零值数、最大跳变和重复整日曲线见 JSON。

四份数据附件数值单元均完整。日期无缺日/重复/倒序。附件 2 两表与附件 4 日期和时段一一对齐。附件 3 的 1095 个空字符串日期单元属于重复日期省写，已按上下文前填；1460 个 6 小时发布点与 35040 个发布—提前期键齐全。9 个输入工作簿在读取前后 SHA-256 相同，原件未修改。

没有零/负电价，因此本附件不会因负价与免费弃电的组合直接引起负价套利无界。未发现数据缺失对 Q1–Q4 构成阻塞。

## 能量与时间验证

- 5000 kW×10/60 h=833.3333333333334 kWh。
- 365×144=52560 个全年区间；2–12 月 334×144=48096 个区间；无跨日重复或空隙。
- Q1 负载电量 {report['totals']['q1_load_kwh']:.10f} kWh，光伏电量 {report['totals']['q1_pv_kwh']:.10f} kWh。
- 读取导出 CSV 后重算功率与电量，最大残差 {max(report['residuals'].values()):.3g}，检查容差 1e-7。
- Q1 周期端点线性积分的总电量与主假设相同，但时段分配变化；必须重新优化才能判断费用敏感性。

## 需保留的解释边界与模板缺陷

1. DQ001：原始功率未标明平均/瞬时。主假设和独立 Q1 敏感性输入见数据字典。
2. DQ002：**五模板均缺当天 00:00–00:10 的正确标签**，从 00:10–00:20 开始。result1 末行是次日 00:00+1–00:10+1，其他宽表末列为混合标签 00:00–00:10+1。保持原标签且将其默认为当天标准 144 时段，会掩盖错位。输出须公开逐格时间映射；正式标签解释仍待官方澄清。缺陷只影响模板语义认证，不妨碍独立标准网格的真实求解。
3. DQ003：充放电/紧急表只有示例日期块，含省略号；Q2/Q4-2 电池表 20 行、Q3/Q4-3 电池表 26 行，紧急表均 11 行。正式输出需在复制件扩展至 334 日，不能误以为只填样例日期就完成。
4. DQ004：预报列为未来 1–24 小时整点，没有 lead 0。插值边界须来自已可见数据/明确的外推，不能偷读未来真实功率。
5. DQ005：附件 4 没有价格发布时点；波动价格是结算真值，不自动成为 0 点已知曲线。
6. DQ006：附件 1 负载与附件 2 的全年逐时段均值最大差仅 {report['cross_source_relationships']['q1_load_vs_annual_slot_mean_max_abs_kw']:.8g} kW；附件 1 电价与附件 4 全年逐时段均值最大差 {report['cross_source_relationships']['q1_price_vs_annual_slot_mean_max_abs']:.8g} 元/kWh，光伏对应最大差 {report['cross_source_relationships']['q1_pv_vs_annual_slot_mean_max_abs_kw']:.8g} kW。这是离线数值关系，不能据此断言命题数据的生成方式。尤其不能用附件 1 的负载/光伏来初始化 Q2–Q4 预测，也不能因此将固定价曲线作为提前知道全年波动价的证据。

结论：物理数值和日期覆盖足以继续建模，时间语义、预报重构及模板标签需要显式假设。不存在从本次数值审计得出的删值依据。
"""
    (root / "context/DATA_QUALITY.md").write_text(quality)
    matrix = """# C 题交付矩阵

依据题面第 1–3 页与五个原模板的实际内容。这里列交付要求；数值完成状态以 CURRENT_STATE 与真实结果验证文件为准。

| 问题 | 允许输入 | 决策时刻 | 必须保存的官方结果 | 全部覆盖 | 论文展示 |
| --- | --- | --- | --- | --- | --- |
| Q1 | 附件 1 电价/负载/光伏预测、附录 1 | 当天 0 点 | result1.xlsx：计划购电量、充放电量 | 当天 144 区间；6 个四小时充放电块及首末状态 | 表 1 六个指定 10 分钟区间+全天量/费；表 2 |
| Q2 | 附件 1 固定电价、附件 2 已观测历史；不可用附件 3 | 每日 0 点定普通合同，日内电池反馈及紧急补缺 | result2.xlsx：计划购电量、充放电量、紧急购电量 | 2025-02-01–12-31，334 日；SOC 跨日连续 | 03-20、06-21、09-23、12-21 的表 1/2/3 |
| Q3 | Q2 输入+截至决策时点已发布附件 3 | 0 点计划；6/12/18 点可改未执行合同 | result3.xlsx：计划购电量、调整购电量、充放电量、紧急购电量 | 同上；内部另存每次调整事件流水 | 同上；分析增加预报/修改机会的价值 |
| Q4-2 | Q2 输入改用附件 4 价格，历史信息边界需声明 | 同 Q2 | result4-2.xlsx：与 result2 对应三表 | 同上 | 波动电价下对应结果 |
| Q4-3 | Q3 输入改用附件 4 价格，历史信息边界需声明 | 同 Q3 | result4-3.xlsx：与 result3 对应四表 | 同上 | 波动电价下对应结果 |

购电量、充放电量、状态均为 kWh；购电费为元。充电/放电分别汇总为交流母线侧输入/输出，不能只给净值。表 1 的指定时段为 10:00–10:10、12:00–12:10、14:00–14:10、16:00–16:10、18:00–18:10、20:00–20:10。表 2 为 0–4、4–8、8–12、12–16、16–20、20–24 小时，另列 0 点/24 点状态。

## 模板结构与使用限制

原模板位于 `data/templates/2026C/`，均为未填数值的空白模板，原件只读。所有原工作簿无合并单元格、无公式。计划和调整宽表各 335×147（含表头、日期、144 槽位和全天量/费）；Q1 购电表 145×2，电池表 7×5。所有原模板的完整表头、固定单元格、尺寸和 SHA-256 已保存到 `results/audit/data_audit.json`。

五模板计划/调整购电时间标签均错位，详见 `DATA_QUALITY.md` 的 DQ002。输出不能静默把值平移来掩盖原始语义冲突。应同时保留标准当天 144 区间的机器明细与逐格导出映射，并清楚标注模板标签的未消歧状态；在既定输出复制件中处理，原件不得改写。

充放电/紧急表是示范块，需扩展到全区间。紧急事件由连续正值 10 分钟区间合并、电量相加，不跨零值、不默认跨日合并。五个输出读回后核对数值类型、时间与日期覆盖、逐项电量、状态和独立费用账本；不能仅以文件存在判定完成。

正式交付位置：`results/submission/`。内部标准明细、验证与假设文件必须可追溯到原始输入哈希、配置、代码和求解状态。
"""
    (root / "context/DELIVERABLE_MATRIX.md").write_text(matrix)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", type=Path, default=Path("configs/2026c.json"))
    run(parser.parse_args().config)
