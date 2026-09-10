import hashlib
import json
import math
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from dispatch import Battery, optimize
from validate_dispatch import check_dispatch


def assert_physics_independently(result, demand, solar, prices, battery, initial, terminal):
    energy = float(initial)
    assert result["soc"][0] == pytest.approx(initial, abs=1e-7)
    all_grid = []
    for t in range(len(demand)):
        grid, charge, discharge, unused = [float(result[k][t]) for k in ["grid", "charge", "discharge", "unused"]]
        assert all(math.isfinite(v) and v >= -1e-7 for v in [grid, charge, discharge, unused])
        assert charge <= battery.power_kw * battery.dt + 1e-7
        assert discharge <= battery.power_kw * battery.dt + 1e-7
        assert not (charge > 1e-7 and discharge > 1e-7)
        delivered = grid + float(solar[t]) + discharge
        consumed = float(demand[t]) + charge + unused
        assert delivered == pytest.approx(consumed, abs=1e-7)
        energy += charge * battery.eta_charge
        energy -= discharge / battery.eta_discharge
        assert battery.minimum - 1e-7 <= energy <= battery.maximum + 1e-7
        assert energy == pytest.approx(result["soc"][t + 1], abs=1e-7)
        all_grid.append(grid * float(prices[t]))
    assert energy == pytest.approx(terminal, abs=1e-7)
    assert math.fsum(all_grid) == pytest.approx(result["solver"]["objective"], abs=1e-7)


def test_constant_price_without_pv_cyclic_battery_cannot_save_cost():
    demand = np.array([100., 200., 300., 400.])
    prices = np.full(4, 1.7)
    result = optimize(demand, np.zeros(4), prices)
    assert result["solver"]["status"] == 0
    assert result["solver"]["objective"] == pytest.approx(1700, abs=1e-7)
    assert sum(result["charge"]) == pytest.approx(0, abs=1e-7)
    assert sum(result["discharge"]) == pytest.approx(0, abs=1e-7)
    assert_physics_independently(result, demand, np.zeros(4), prices, Battery(), 6000, 6000)


def test_zero_load_zero_pv_requires_no_energy_purchase():
    demand = np.zeros(6)
    prices = np.array([.1, 1., .3, 2., .4, .5])
    result = optimize(demand, demand, prices)
    assert result["solver"]["objective"] == pytest.approx(0, abs=1e-7)
    assert sum(result["grid"]) == pytest.approx(0, abs=1e-7)
    assert result["soc"] == pytest.approx(np.full(7, 6000), abs=1e-7)
    assert_physics_independently(result, demand, demand, prices, Battery(), 6000, 6000)


def test_one_interval_pv_surplus_charges_only_available_headroom():
    battery = Battery(minimum=0, maximum=100, power_kw=600)
    result = optimize([0], [100], [1], initial=90, terminal=100, battery=battery)
    expected_charge = (100 - 90) / .9
    assert result["grid"][0] == pytest.approx(0, abs=1e-7)
    assert result["charge"][0] == pytest.approx(expected_charge, abs=1e-7)
    assert result["unused"][0] == pytest.approx(100 - expected_charge, abs=1e-7)
    assert_physics_independently(result, [0], [100], [1], battery, 90, 100)


def test_analytic_two_interval_arbitrage_binds_power_limit():
    battery = Battery(minimum=0, maximum=100, power_kw=60)
    result = optimize([0, 10], [0, 0], [1, 3], initial=0, terminal=0, battery=battery)
    assert result["charge"] == pytest.approx([10, 0], abs=1e-7)
    assert result["discharge"] == pytest.approx([0, 8.1], abs=1e-7)
    assert result["grid"] == pytest.approx([10, 1.9], abs=1e-7)
    assert result["solver"]["objective"] == pytest.approx(15.7, abs=1e-7)
    assert_physics_independently(result, [0, 10], [0, 0], [1, 3], battery, 0, 0)


@pytest.mark.parametrize("demand,solar,prices", [
    ([0, 10], [0, 0], [1, 3]),
    ([7, 1, 8], [0, 14, 0], [3, 1, 4]),
    ([0, 0, 0], [20, 0, 20], [1, 2, 3]),
])
def test_lp_is_a_lower_bound_under_identical_constraints(demand, solar, prices):
    battery = Battery(minimum=0, maximum=100, power_kw=60)
    integer = optimize(demand, solar, prices, initial=20, terminal=20, battery=battery)
    relaxed = optimize(demand, solar, prices, initial=20, terminal=20, battery=battery, relax=True)
    assert relaxed["solver"]["objective"] <= integer["solver"]["objective"] + 1e-7
    assert_physics_independently(integer, demand, solar, prices, battery, 20, 20)


def test_production_checker_rejects_perturbed_soc():
    result = optimize([100, 100], [0, 0], [1, 1])
    result["soc"][1] += 10
    assert not check_dispatch(np.array([100, 100]), np.zeros(2), np.ones(2), result)["passed"]
    with pytest.raises(AssertionError):
        assert_physics_independently(result, [100, 100], [0, 0], [1, 1], Battery(), 6000, 6000)


def test_original_144_right_endpoints_have_correct_interval_energy():
    source = openpyxl.load_workbook(ROOT / "data/raw/2026C/附件1.xlsx", read_only=True, data_only=True)
    raw = list(source.active.values)[1:]
    derived = pd.read_csv(ROOT / "data/processed/2026C/q1.csv")
    assert len(raw) == len(derived) == 144
    assert raw[0][0].hour == 0 and raw[0][0].minute == 10
    assert raw[-1][0] == "0:00+1"
    assert derived.interval_index.tolist() == list(range(144))
    assert derived.start_minute.tolist() == list(range(0, 1440, 10))
    assert derived.end_minute.tolist() == list(range(10, 1441, 10))
    for (_, raw_price, raw_load, raw_pv), row in zip(raw, derived.itertuples()):
        assert row.price == raw_price
        assert row.load_kwh == pytest.approx(raw_load / 6, abs=1e-10)
        assert row.pv_kwh == pytest.approx(raw_pv / 6, abs=1e-10)
    assert 5000 * (10 / 60) == pytest.approx(833.3333333333334)
    source.close()


@pytest.mark.parametrize("strategy", ["no_storage", "simple_rule", "milp", "lp_relaxation"])
def test_q1_saved_dispatch_reconciles_independently(strategy):
    source = pd.read_csv(ROOT / "data/processed/2026C/q1.csv")
    rows = pd.read_csv(ROOT / f"results/q1/{strategy}_dispatch.csv")
    summary = json.loads((ROOT / "results/q1/summary.json").read_text())
    recorded = summary["strategies"][strategy]
    assert len(rows) == 144
    for field in ["load_kwh", "pv_kwh", "price"]:
        assert rows[field].to_numpy() == pytest.approx(source[field].to_numpy(), abs=1e-10)
    assert rows.soc_start_kwh.iloc[1:].to_numpy() == pytest.approx(rows.soc_end_kwh.iloc[:-1].to_numpy(), abs=1e-8)
    result = {k: rows[f"{k}_kwh"].to_numpy() for k in ["grid", "charge", "discharge", "unused"]}
    result["soc"] = np.r_[rows.soc_start_kwh.iloc[0], rows.soc_end_kwh]
    result["solver"] = recorded["solver"]
    assert_physics_independently(result, source.load_kwh, source.pv_kwh, source.price, Battery(), 6000, 6000)
    for field in ["grid", "charge", "discharge", "unused"]:
        assert math.fsum(result[field]) == pytest.approx(recorded[f"{field}_kwh"], abs=1e-7)
    loss = math.fsum(result["charge"]) - math.fsum(result["discharge"])
    supplied_minus_load = math.fsum(result["grid"]) + math.fsum(source.pv_kwh) - math.fsum(source.load_kwh)
    assert loss + math.fsum(result["unused"]) == pytest.approx(supplied_minus_load, abs=1e-7)
    assert math.fsum(result["discharge"]) == pytest.approx(.81 * math.fsum(result["charge"]), abs=1e-7)


def test_q1_tables_match_six_named_intervals_and_four_hour_blocks():
    rows = pd.read_csv(ROOT / "results/q1/milp_dispatch.csv")
    table1 = pd.read_csv(ROOT / "results/q1/table1.csv")
    table2 = pd.read_csv(ROOT / "results/q1/table2.csv")
    assert table1.start_minute.tolist() == [600, 720, 840, 960, 1080, 1200]
    assert table2.start_hour.tolist() == [0, 4, 8, 12, 16, 20]
    for r in table1.itertuples():
        source = rows[rows.start_minute == r.start_minute].iloc[0]
        assert r.end_minute == r.start_minute + 10
        assert r.grid_kwh == pytest.approx(source.grid_kwh, abs=1e-9)
    for r in table2.itertuples():
        block = rows[(rows.start_minute >= r.start_hour * 60) & (rows.end_minute <= r.end_hour * 60)]
        assert len(block) == 24
        assert r.charge_kwh == pytest.approx(math.fsum(block.charge_kwh), abs=1e-7)
        assert r.discharge_kwh == pytest.approx(math.fsum(block.discharge_kwh), abs=1e-7)


def test_q1_input_and_code_fingerprints_match_current_files():
    summary = json.loads((ROOT / "results/q1/summary.json").read_text())
    digest = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
    assert summary["input_sha256"] == digest(ROOT / "data/processed/2026C/q1.csv")
    for name, expected in summary["code_sha256"].items():
        assert digest(ROOT / "src" / name) == expected


@pytest.mark.parametrize("case", ["round_trip_90_percent", "instantaneous_points_periodic_linear"])
def test_q1_sensitivity_reproduces_saved_cost_and_independent_physics(case):
    source = pd.read_csv(ROOT / "data/processed/2026C/q1.csv")
    battery = Battery()
    if case == "round_trip_90_percent":
        battery = Battery(eta_charge=math.sqrt(.9), eta_discharge=math.sqrt(.9))
    else:
        source = pd.read_csv(ROOT / "data/processed/2026C/q1_instantaneous_sensitivity.csv")
    result = optimize(source.load_kwh, source.pv_kwh, source.price, battery=battery)
    summary = json.loads((ROOT / "results/q1/summary.json").read_text())
    recorded = next(r for r in summary["sensitivity"] if r["case"] == case)
    assert result["solver"]["objective"] == pytest.approx(recorded["cost_yuan"], abs=1e-6)
    assert math.fsum(result["grid"]) == pytest.approx(recorded["grid_kwh"], abs=1e-6)
    assert_physics_independently(result, source.load_kwh, source.pv_kwh, source.price, battery, 6000, 6000)


def test_q1_template_defect_is_disclosed_not_silently_certified():
    original = openpyxl.load_workbook(ROOT / "data/templates/2026C/result1.xlsx", data_only=True)
    filled = openpyxl.load_workbook(ROOT / "results/submission/result1.xlsx", data_only=True)
    rows = pd.read_csv(ROOT / "results/q1/milp_dispatch.csv")
    assert original["计划购电量"]["A2"].value == "0:10-0:20"
    assert original["计划购电量"]["A145"].value == "0:00+1-0:10+1"
    assert filled.sheetnames == original.sheetnames
    for sheet in original:
        for source_row in sheet:
            for cell in source_row:
                if cell.value is not None:
                    assert filled[sheet.title][cell.coordinate].value == cell.value
    for i in range(144):
        assert filled["计划购电量"].cell(i + 2, 1).value == original["计划购电量"].cell(i + 2, 1).value
        assert filled["计划购电量"].cell(i + 2, 2).value == pytest.approx(rows.grid_kwh.iloc[i], abs=1e-7)
    for block in range(6):
        source_rows = rows.iloc[24 * block:24 * (block + 1)]
        assert filled["充放电量"].cell(block + 2, 2).value == pytest.approx(math.fsum(source_rows.charge_kwh), abs=1e-7)
        assert filled["充放电量"].cell(block + 2, 3).value == pytest.approx(math.fsum(source_rows.discharge_kwh), abs=1e-7)
    assert filled["充放电量"]["E2"].value == pytest.approx(rows.soc_start_kwh.iloc[0], abs=1e-7)
    assert filled["充放电量"]["E3"].value == pytest.approx(rows.soc_end_kwh.iloc[-1], abs=1e-7)
    summary = json.loads((ROOT / "results/q1/summary.json").read_text())
    assert summary["template_status"] == "TEMPLATE_TIME_LABEL_AMBIGUITY"
    original.close()
    filled.close()
